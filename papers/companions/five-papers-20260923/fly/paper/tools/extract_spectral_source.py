"""Read-only workbook extraction; keep source curves, cells and units explicit."""
from pathlib import Path
import hashlib
import json
import math
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
APP=ROOT/'application/spectral-receiver-bridge-v0'


def main():
    plan=json.loads((APP/'ANALYSIS-PLAN.json').read_text('utf-8'))
    source=ROOT/plan['source']['path']
    assert hashlib.sha256(source.read_bytes()).hexdigest()==plan['source']['sha256']
    w=openpyxl.load_workbook(source,read_only=True,data_only=False)
    assert w.sheetnames==['Sheet 1']
    s=w['Sheet 1']
    specs=[('Rh1-low','Rh1','D','E','F',315,550),
           ('Rh3-low','Rh3','D','G','H',315,550),
           ('Rh4-low','Rh4','D','I','J',315,550),
           ('Rh5-low','Rh5','D','K','L',315,550),
           ('Rh6-low','Rh6','D','M','N',315,550),
           ('Rh1-high','Rh1','O','P','Q',450,700),
           ('Rh6-high','Rh6','O','R','S',450,700)]
    curves=[]
    for key,receptor,x,m,sd,first,last in specs:
        assert s[m+'1'].value==receptor
        assert s[x+'2'].value=='Wavelength (nm)' and s[m+'2'].value=='Mean'
        assert s[sd+'2'].value in ('Standard devation','Standard deviation')
        values=[]
        for row,lam in enumerate(range(first,last+1,5),start=3):
            cells={k:v+str(row) for k,v in [('wavelength',x),('mean',m),('sd',sd)]}
            a,b,c=[s[cells[k]].value for k in ('wavelength','mean','sd')]
            assert a==lam and all(isinstance(v,(int,float)) and math.isfinite(v) for v in (b,c))
            assert b>=0 and c>=0
            values.append(dict(wavelengthNm=a,meanMv=b,sdMv=c,cells=cells))
        curves.append(dict(id=key,receptor=receptor,wavelengthRangeNm=[first,last],
                           reportedAnimalsPerCurve=6,points=values))
    assert sum(len(c['points']) for c in curves)==342
    outside=[c.coordinate for row in s.iter_rows() for c in row
             if c.value is not None and not 4<=c.column<=19]
    assert not outside
    w.close()
    result=dict(source=plan['source'],sheet='Sheet 1',curves=curves,
                unitsSource='Supplementary Figure S6, PDF page 7: photoreceptor response (mV)',
                dataLevel='published processed curve means and between-animal SD, not raw individual traces',
                rawAnimalRecordsAvailable=False,sourceFormulasEvaluated=False,
                extractorSha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest())
    out=APP/'source-extraction-01'
    out.mkdir(exist_ok=False)
    with (out/'CURVES.json').open('x',encoding='utf-8') as f: json.dump(result,f,indent=2)
    print(json.dumps({'curves':len(curves),'meanSdPairs':342,'output':str(out/'CURVES.json')}))


if __name__=='__main__': main()
